# -*- coding: utf-8 -*-
"""
formato_docx.py — la identidad visual de los documentos del proyecto.

Los .docx no se editan a mano: se generan con un script, para poder
rehacerlos con datos nuevos sin perder el formato. Este módulo tiene el
formato; cada documento pone solo su contenido.

La paleta sale del propio sitio: el gris de las tablas es el mismo
--card del modo claro.

Trampas que este módulo ya resuelve, y que costaron descubrir:
  - fijar el ancho de una celda no basta. Word y LibreOffice usan la
    rejilla de la tabla, y sin escribir el w:tblGrid todas las columnas
    salen iguales.
  - restar dos Length de python-docx devuelve un int en EMU, que ya no
    tiene .twips.
  - no existe negrita a mitad de párrafo: hay que partirlo en corridas.
    Para eso está pm().
"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

TINTA = RGBColor(0x11, 0x11, 0x11)
SUAVE = RGBColor(0x5A, 0x5A, 0x5A)
ACENTO = RGBColor(0x1F, 0x4E, 0x5F)
ACENTO_HEX = "1F4E5F"
ALERTA_HEX = "9B2C1F"
FONDO_HEX = "F2F1EF"
LINEA_HEX = "D8D6D2"
FUENTE = "Segoe UI"


# ---------------------------------------------------------- XML a mano

def _sombrear(celda, hexc):
    e = OxmlElement("w:shd")
    e.set(qn("w:val"), "clear"); e.set(qn("w:color"), "auto"); e.set(qn("w:fill"), hexc)
    celda._tc.get_or_add_tcPr().append(e)


def _bordes(celda, arriba=None, abajo=None, izq=None, der=None):
    marco = OxmlElement("w:tcBorders")
    for nombre, spec in (("top", arriba), ("bottom", abajo), ("left", izq), ("right", der)):
        b = OxmlElement("w:" + nombre)
        if spec:
            b.set(qn("w:val"), "single"); b.set(qn("w:sz"), str(spec[0])); b.set(qn("w:color"), spec[1])
        else:
            b.set(qn("w:val"), "nil")
        marco.append(b)
    celda._tc.get_or_add_tcPr().append(marco)


def _margenes(celda, a=70, b=70, i=120, d=120):
    m = OxmlElement("w:tcMar")
    for nombre, valor in (("top", a), ("bottom", b), ("left", i), ("right", d)):
        el = OxmlElement("w:" + nombre); el.set(qn("w:w"), str(valor)); el.set(qn("w:type"), "dxa")
        m.append(el)
    celda._tc.get_or_add_tcPr().append(m)


def _borde_inferior(par, grosor, color):
    marco = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single"); b.set(qn("w:sz"), str(grosor))
    b.set(qn("w:space"), "6"); b.set(qn("w:color"), color)
    marco.append(b)
    par._p.get_or_add_pPr().append(marco)


def _tracking(run, v):
    e = OxmlElement("w:spacing"); e.set(qn("w:val"), str(v))
    run._element.get_or_add_rPr().append(e)


# ---------------------------------------------------------- el documento

class Doc(object):
    def __init__(self, titulo, autor="Vicente Cáceres Farías", margen=1.9):
        self.doc = Document()
        self.doc.core_properties.title = titulo
        self.doc.core_properties.author = autor
        sec = self.doc.sections[0]
        for lado in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
            setattr(sec, lado, Cm(margen))
        self.ancho = sec.page_width - sec.left_margin - sec.right_margin

        n = self.doc.styles["Normal"]
        n.font.name = FUENTE
        n.font.size = Pt(10)
        n.font.color.rgb = TINTA
        n._element.rPr.rFonts.set(qn("w:eastAsia"), FUENTE)
        n.paragraph_format.space_after = Pt(6)
        n.paragraph_format.line_spacing = 1.20

    # ---- párrafos
    def p(self, texto="", size=10, color=TINTA, bold=False, antes=0, despues=6,
          align=None, caps=False, track=None):
        par = self.doc.add_paragraph()
        par.paragraph_format.space_before = Pt(antes)
        par.paragraph_format.space_after = Pt(despues)
        if align is not None:
            par.alignment = align
        r = par.add_run(texto)
        r.font.name = FUENTE; r.font.size = Pt(size); r.font.color.rgb = color
        r.bold = bold; r.font.all_caps = caps
        if track:
            _tracking(r, track)
        return par

    def pm(self, texto, antes=0, despues=6, size=10):
        """Como p(), pero los tramos entre ** salen en negrita de verdad."""
        par = self.doc.add_paragraph()
        par.paragraph_format.space_before = Pt(antes)
        par.paragraph_format.space_after = Pt(despues)
        for i, trozo in enumerate(texto.split("**")):
            if not trozo:
                continue
            r = par.add_run(trozo)
            r.font.name = FUENTE; r.font.size = Pt(size); r.font.color.rgb = TINTA
            r.bold = (i % 2 == 1)
        return par

    def h1(self, texto):
        par = self.p(texto, size=15, bold=True, antes=20, despues=3)
        _borde_inferior(par, 12, ACENTO_HEX)
        return par

    def h2(self, texto):
        return self.p(texto, size=11.5, bold=True, color=ACENTO, antes=13, despues=4)

    def etiqueta(self, texto):
        return self.p(texto, size=7.5, bold=True, color=SUAVE, caps=True,
                      track=40, antes=10, despues=2)

    def vinieta(self, texto, negrita=None):
        par = self.doc.add_paragraph()
        par.paragraph_format.left_indent = Cm(0.62)
        par.paragraph_format.first_line_indent = Cm(-0.38)
        par.paragraph_format.space_after = Pt(3.5)
        g = par.add_run("—  ")
        g.font.name = FUENTE; g.font.size = Pt(10); g.font.color.rgb = ACENTO
        if negrita:
            rb = par.add_run(negrita)
            rb.font.name = FUENTE; rb.font.size = Pt(10); rb.bold = True; rb.font.color.rgb = TINTA
        r = par.add_run(texto)
        r.font.name = FUENTE; r.font.size = Pt(10); r.font.color.rgb = TINTA
        return par

    def paso(self, n, titulo, texto):
        par = self.doc.add_paragraph()
        par.paragraph_format.left_indent = Cm(0.95)
        par.paragraph_format.first_line_indent = Cm(-0.95)
        par.paragraph_format.space_before = Pt(8)
        par.paragraph_format.space_after = Pt(2)
        rn = par.add_run("%s  " % n)
        rn.font.name = FUENTE; rn.font.size = Pt(11); rn.bold = True; rn.font.color.rgb = ACENTO
        rt = par.add_run(titulo)
        rt.font.name = FUENTE; rt.font.size = Pt(11); rt.bold = True; rt.font.color.rgb = TINTA
        p2 = self.doc.add_paragraph()
        p2.paragraph_format.left_indent = Cm(0.95)
        p2.paragraph_format.space_after = Pt(6)
        r2 = p2.add_run(texto)
        r2.font.name = FUENTE; r2.font.size = Pt(10); r2.font.color.rgb = TINTA

    def salto(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---- tablas
    def _fijar_anchos(self, t, anchos):
        tbl = t._tbl
        lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed")
        tbl.tblPr.append(lay)
        vieja = tbl.find(qn("w:tblGrid"))
        if vieja is not None:
            tbl.remove(vieja)
        grid = OxmlElement("w:tblGrid")
        for a in anchos:
            gc = OxmlElement("w:gridCol")
            tw = a.twips if hasattr(a, "twips") else int(a) // 635
            gc.set(qn("w:w"), str(int(tw))); grid.append(gc)
        tbl.tblPr.addnext(grid)
        for fila in t.rows:
            for c, a in zip(fila.cells, anchos):
                c.width = a

    def tabla(self, filas, anchos, encabezado=True):
        t = self.doc.add_table(rows=0, cols=len(anchos))
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        t.autofit = False
        for i, fila in enumerate(filas):
            celdas = t.add_row().cells
            for j, texto in enumerate(fila):
                c = celdas[j]
                _margenes(c)
                _bordes(c, arriba=(4, LINEA_HEX), abajo=(4, LINEA_HEX))
                if encabezado and i == 0:
                    _sombrear(c, FONDO_HEX)
                par = c.paragraphs[0]
                par.paragraph_format.space_after = Pt(0)
                par.paragraph_format.space_before = Pt(0)
                par.paragraph_format.line_spacing = 1.12
                r = par.add_run(texto)
                r.font.name = FUENTE; r.font.size = Pt(9)
                es_tit = encabezado and i == 0
                r.bold = es_tit
                r.font.color.rgb = TINTA if (es_tit or j == 0) else SUAVE
        self._fijar_anchos(t, anchos)
        self.p(despues=4)
        return t

    def lineas_para_escribir(self, campos, ancho_etq=Cm(3.6), alto=Cm(0.95)):
        """Renglones en blanco con su etiqueta, para llenar a mano."""
        t = self.doc.add_table(rows=0, cols=2)
        t.autofit = False
        anchos = [ancho_etq, self.ancho - ancho_etq]
        for etq, n in campos:
            for k in range(n):
                fila = t.add_row()
                fila.height = alto
                fila.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
                for c in fila.cells:
                    _margenes(c, 110, 110, 120, 120)
                    _bordes(c, abajo=(4, LINEA_HEX))
                par = fila.cells[0].paragraphs[0]
                par.paragraph_format.space_after = Pt(0)
                if k == 0:
                    r = par.add_run(etq)
                    r.font.name = FUENTE; r.font.size = Pt(9); r.bold = True; r.font.color.rgb = TINTA
                fila.cells[1].paragraphs[0].paragraph_format.space_after = Pt(0)
        self._fijar_anchos(t, anchos)
        self.p(despues=4)
        return t

    def recuadro(self, titulo, texto, barra=ACENTO_HEX):
        t = self.doc.add_table(rows=1, cols=1)
        t.autofit = False
        c = t.rows[0].cells[0]
        _sombrear(c, FONDO_HEX)
        _margenes(c, 160, 160, 200, 200)
        _bordes(c, izq=(18, barra))
        par = c.paragraphs[0]; par.paragraph_format.space_after = Pt(3)
        r = par.add_run(titulo)
        r.font.name = FUENTE; r.font.size = Pt(10); r.bold = True; r.font.color.rgb = TINTA
        p2 = c.add_paragraph(); p2.paragraph_format.space_after = Pt(0)
        p2.paragraph_format.line_spacing = 1.18
        r2 = p2.add_run(texto)
        r2.font.name = FUENTE; r2.font.size = Pt(9.5); r2.font.color.rgb = SUAVE
        self._fijar_anchos(t, [self.ancho])
        self.p(despues=6)
        return t

    # ---- portada y cierre
    def portada(self, sobretitulo, titulo, bajada, fecha):
        self.p(sobretitulo, size=8, bold=True, color=SUAVE, caps=True, track=120, despues=2)
        par = self.doc.add_paragraph(); par.paragraph_format.space_after = Pt(2)
        r = par.add_run(titulo)
        r.font.name = FUENTE; r.font.size = Pt(26); r.bold = True; r.font.color.rgb = TINTA
        par = self.doc.add_paragraph(); par.paragraph_format.space_after = Pt(9)
        r = par.add_run(bajada)
        r.font.name = FUENTE; r.font.size = Pt(13); r.font.color.rgb = SUAVE
        _borde_inferior(par, 18, ACENTO_HEX)
        self.p(fecha, size=10, bold=True, color=ACENTO, antes=8, despues=14)

    def pie(self, texto):
        par = self.doc.add_paragraph()
        par.paragraph_format.space_before = Pt(16)
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = par.add_run(texto)
        r.font.name = FUENTE; r.font.size = Pt(8); r.font.color.rgb = SUAVE

    def guardar(self, ruta):
        self.doc.save(ruta)
        print("escrito:", ruta)
